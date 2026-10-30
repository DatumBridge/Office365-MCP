"""Read .xlsx and .xlsm cell values into JSON or XML.

Formulas are returned as formula text and are not evaluated. Macros are not
executed. Legacy BIFF .xls (OLE compound file) is rejected.
"""

from __future__ import annotations

import io
import json
import logging
import math
import re
import zipfile
from datetime import date, datetime, time
from pathlib import PurePosixPath
from typing import Any, Optional
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter

from app.core.exceptions import GraphError, GraphValidationError

logger = logging.getLogger(__name__)

MAX_WORKBOOK_BYTES = 8 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 32 * 1024 * 1024
MAX_ZIP_ENTRIES = 512
MAX_SHEETS = 32
MAX_ROWS = 2000
MAX_COLUMNS = 50
MAX_CELLS = 20000
MAX_CELL_CHARS = 8000
MAX_OUTPUT_CHARS = 1_000_000

_EXCEL_EXTENSIONS = frozenset({".xlsx", ".xlsm", ".xls"})
_EXCEL_MIMES = frozenset(
    {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel.sheet.macroenabled.12",
        "application/vnd.ms-excel",
    }
)
_ALLOWED_ZIP = frozenset({zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED})
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
_EXCEL_ROW_KEY = "excel_row"
_FORMATS = frozenset({"json", "xml"})


def is_excel_file(file_name: Optional[str], mime_type: Optional[str]) -> bool:
    """True when the drive item should be parsed as a workbook."""
    ext = PurePosixPath(file_name or "").suffix.lower()
    if ext in {".xlsb"}:
        return False
    if ext in _EXCEL_EXTENSIONS:
        return True
    mime = (mime_type or "").strip().lower()
    return mime in _EXCEL_MIMES


def normalize_excel_format(excel_format: Optional[str]) -> Optional[str]:
    """Return json or xml, or None when the caller did not ask for a transform."""
    if excel_format is None:
        return None
    text = str(excel_format).strip().lower()
    if not text:
        return None
    if text not in _FORMATS:
        raise GraphValidationError("excel_format must be json or xml")
    return text


def transform_excel(
    content: bytes,
    *,
    excel_format: str,
    file_name: Optional[str] = None,
    sheet_name: Optional[str] = None,
    header_row: bool = False,
    max_bytes: int = MAX_WORKBOOK_BYTES,
    max_uncompressed_bytes: int = MAX_UNCOMPRESSED_BYTES,
    max_zip_entries: int = MAX_ZIP_ENTRIES,
    max_sheets: int = MAX_SHEETS,
    max_rows: int = MAX_ROWS,
    max_columns: int = MAX_COLUMNS,
    max_cells: int = MAX_CELLS,
    max_cell_chars: int = MAX_CELL_CHARS,
    max_output_chars: int = MAX_OUTPUT_CHARS,
) -> dict[str, Any]:
    """Parse workbook bytes into a JSON document and/or an XML document.

    The returned dict always includes format, truncated, and content_text.
    content_json is set for json. content_xml is set for xml.
    """
    fmt = normalize_excel_format(excel_format)
    if fmt is None:
        raise GraphValidationError("excel_format must be json or xml")
    if not isinstance(content, (bytes, bytearray)) or not content:
        raise GraphValidationError("Could not read Excel workbook")
    _require_defusedxml()

    workbook = _load_workbook(
        bytes(content),
        max_bytes=max_bytes,
        max_uncompressed_bytes=max_uncompressed_bytes,
        max_zip_entries=max_zip_entries,
    )
    try:
        try:
            sheets, truncated = _read_sheets(
                workbook,
                sheet_name=_clean_sheet_name(sheet_name),
                header_row=bool(header_row),
                max_sheets=max_sheets,
                max_rows=max_rows,
                max_columns=max_columns,
                max_cells=max_cells,
                max_cell_chars=max_cell_chars,
                max_output_chars=max_output_chars,
            )
        except GraphError:
            raise
        except Exception as exc:
            logger.warning("excel read failed: %s", type(exc).__name__)
            raise GraphValidationError("Could not read Excel workbook") from exc
    finally:
        workbook.close()

    document = {
        "file_name": file_name or "",
        "truncated": truncated,
        "sheets": sheets,
    }
    text, truncated = _shrink_to_fit(document, fmt, max_output_chars)
    document["truncated"] = truncated
    if fmt == "json":
        return {
            "format": "json",
            "truncated": truncated,
            "content_json": document,
            "content_xml": None,
            "content_text": text,
        }
    return {
        "format": "xml",
        "truncated": truncated,
        "content_json": None,
        "content_xml": text,
        "content_text": text,
    }


def _clean_sheet_name(sheet_name: Optional[str]) -> Optional[str]:
    if sheet_name is None:
        return None
    text = str(sheet_name).strip()
    return text or None


def _require_defusedxml() -> None:
    import openpyxl
    from openpyxl.xml import functions as xml_functions

    if not getattr(xml_functions, "DEFUSEDXML", False):
        raise GraphValidationError("Excel parser is not available")
    # Touch the public flag when openpyxl exports it. Older builds only set
    # the functions module flag, which is the one the parser actually uses.
    getattr(openpyxl, "DEFUSEDXML", None)


def _render(document: dict[str, Any], fmt: str) -> str:
    if fmt == "json":
        return _json_text(document)
    return _to_xml(document)


def _shrink_to_fit(
    document: dict[str, Any],
    fmt: str,
    limit: int,
) -> tuple[str, bool]:
    """Keep the serialized workbook within limit, including repeated header keys."""
    truncated = bool(document.get("truncated"))
    text = _render(document, fmt)
    if len(text) <= limit:
        return text, truncated
    truncated = True
    document["truncated"] = True
    while len(text) > limit and _drop_last_row(document):
        text = _render(document, fmt)
    if len(text) <= limit:
        return text, True
    clipped = _clip_strings(document, max(1, limit // 40))
    if isinstance(clipped, dict):
        document.clear()
        document.update(clipped)
    text = _render(document, fmt)
    if len(text) <= limit:
        return text, True
    document["sheets"] = []
    return _render(document, fmt), True


def _drop_last_row(document: dict[str, Any]) -> bool:
    sheets = document.get("sheets") or []
    for sheet in reversed(sheets):
        rows = sheet.get("rows")
        if isinstance(rows, list) and rows:
            rows.pop()
            return True
    if len(sheets) > 1:
        sheets.pop()
        return True
    return False


def _clip_strings(value: Any, limit: int) -> Any:
    if isinstance(value, str):
        return value[:limit]
    if isinstance(value, list):
        return [_clip_strings(item, limit) for item in value]
    if isinstance(value, dict):
        clipped = {}
        for key, item in value.items():
            clipped[_clip_strings(key, limit) if isinstance(key, str) else key] = (
                _clip_strings(item, limit)
            )
        return clipped
    return value


def _json_text(document: dict[str, Any]) -> str:
    return json.dumps(document, ensure_ascii=False, separators=(",", ":"))


def _load_workbook(
    content: bytes,
    *,
    max_bytes: int,
    max_uncompressed_bytes: int,
    max_zip_entries: int,
):
    if len(content) > max_bytes:
        raise GraphValidationError(
            f"Excel workbook exceeds the {max_bytes} byte read limit"
        )
    if content.startswith(_OLE_MAGIC):
        raise GraphValidationError(
            "Legacy .xls workbooks cannot be transformed. "
            "Save the file as .xlsx and read it again."
        )
    if not content.startswith(b"PK"):
        raise GraphValidationError("Could not read Excel workbook")
    safe = _safe_office_bytes(
        content,
        max_uncompressed_bytes=max_uncompressed_bytes,
        max_zip_entries=max_zip_entries,
    )
    try:
        return load_workbook(
            io.BytesIO(safe),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except GraphError:
        raise
    except Exception as exc:
        logger.warning("excel open failed: %s", type(exc).__name__)
        raise GraphValidationError("Could not read Excel workbook") from exc


def _safe_office_bytes(
    content: bytes,
    *,
    max_uncompressed_bytes: int,
    max_zip_entries: int,
) -> bytes:
    """Copy a workbook into a stored ZIP, counting inflated bytes.

    Only store and deflate members are accepted. The copy is what openpyxl
    parses, so a declared file_size cannot hide a larger inflate.
    """
    try:
        package = zipfile.ZipFile(io.BytesIO(content))
    except zipfile.BadZipFile as exc:
        raise GraphValidationError("Could not read Excel workbook") from exc

    try:
        infos = package.infolist()
        if len(infos) > max_zip_entries:
            raise GraphValidationError("Excel workbook exceeds the entry limit")

        safe = io.BytesIO()
        total = 0
        has_workbook = False
        with zipfile.ZipFile(safe, "w", compression=zipfile.ZIP_STORED) as stored:
            for info in infos:
                name = info.filename.replace("\\", "/")
                parts = [part for part in name.split("/") if part not in ("", ".")]
                if name.startswith("/") or ".." in parts:
                    raise GraphValidationError("Could not read Excel workbook")
                lowered = "/".join(parts).lower()
                if lowered in ("encryptioninfo", "encryptedpackage"):
                    raise GraphValidationError("Encrypted Excel workbooks cannot be read")
                if info.compress_type not in _ALLOWED_ZIP:
                    raise GraphValidationError("Could not read Excel workbook")
                if lowered == "xl/workbook.xml":
                    has_workbook = True
                member = _read_member(package, info, max_uncompressed_bytes - total)
                total += len(member)
                if total > max_uncompressed_bytes:
                    raise GraphValidationError(
                        "Excel workbook exceeds the expanded size limit"
                    )
                stored.writestr(name, member)
        if not has_workbook:
            raise GraphValidationError(
                "Only .xlsx and .xlsm workbooks can be transformed to JSON or XML"
            )
        return safe.getvalue()
    finally:
        package.close()


def _read_member(package: zipfile.ZipFile, info: zipfile.ZipInfo, room: int) -> bytes:
    if room < 0:
        raise GraphValidationError("Excel workbook exceeds the expanded size limit")
    chunks = []
    seen = 0
    try:
        handle = package.open(info, "r")
    except RuntimeError as exc:
        raise GraphValidationError("Could not read Excel workbook") from exc
    try:
        while True:
            block = handle.read(64 * 1024)
            if not block:
                break
            seen += len(block)
            if seen > room:
                raise GraphValidationError(
                    "Excel workbook exceeds the expanded size limit"
                )
            chunks.append(block)
    finally:
        handle.close()
    return b"".join(chunks)


def _read_sheets(
    workbook,
    *,
    sheet_name: Optional[str],
    header_row: bool,
    max_sheets: int,
    max_rows: int,
    max_columns: int,
    max_cells: int,
    max_cell_chars: int,
    max_output_chars: int,
) -> tuple[list[dict[str, Any]], bool]:
    selected, truncated = _select_worksheets(workbook, sheet_name, max_sheets)
    sheets: list[dict[str, Any]] = []
    cells_left = max_cells
    chars_left = max_output_chars
    for worksheet in selected:
        parsed, sheet_truncated, cells_left, chars_left = _read_worksheet(
            worksheet,
            max_rows=max_rows,
            max_columns=max_columns,
            cells_left=cells_left,
            chars_left=chars_left,
            max_cell_chars=max_cell_chars,
        )
        truncated = truncated or sheet_truncated
        sheets.append(_sheet_document(worksheet.title, parsed, header_row))
        if (cells_left <= 0 or chars_left <= 0) and worksheet is not selected[-1]:
            truncated = True
            break
    return sheets, truncated


def _select_worksheets(workbook, sheet_name: Optional[str], max_sheets: int):
    worksheets = list(workbook.worksheets)
    if sheet_name is None:
        truncated = len(worksheets) > max_sheets
        return worksheets[:max_sheets], truncated

    exact = [ws for ws in worksheets if ws.title == sheet_name]
    if len(exact) == 1:
        return exact, False
    folded = sheet_name.casefold()
    matches = [ws for ws in worksheets if ws.title.casefold() == folded]
    if len(matches) == 1:
        return matches, False
    if len(matches) > 1:
        raise GraphValidationError(f"Worksheet name '{sheet_name}' is ambiguous")
    available = ", ".join(ws.title for ws in worksheets[:max_sheets])
    raise GraphValidationError(
        f"Worksheet '{sheet_name}' was not found. Sheets: {available}"
    )


def _read_worksheet(
    worksheet,
    *,
    max_rows: int,
    max_columns: int,
    cells_left: int,
    chars_left: int,
    max_cell_chars: int,
) -> tuple[list[tuple[int, list[Any]]], bool, int, int]:
    if not hasattr(worksheet, "iter_rows"):
        return [], False, cells_left, chars_left

    truncated = False
    parsed: list[tuple[int, list[Any]]] = []
    try:
        row_iter = worksheet.iter_rows(values_only=False)
    except (AttributeError, TypeError, ValueError):
        return [], False, cells_left, chars_left

    for cells in row_iter:
        if not cells:
            continue
        excel_row = _excel_row(cells)
        if excel_row is None:
            continue
        if excel_row > max_rows:
            if _row_has_value(cells, max_cell_chars):
                truncated = True
                break
            continue
        values, row_truncated, cells_left, chars_left, extra_column = _row_values(
            cells,
            max_columns=max_columns,
            cells_left=cells_left,
            chars_left=chars_left,
            max_cell_chars=max_cell_chars,
        )
        truncated = truncated or row_truncated or extra_column
        if any(value is not None for value in values):
            parsed.append((excel_row, _trim_trailing(values)))
        if row_truncated:
            break
    return parsed, truncated, cells_left, chars_left


def _excel_row(cells) -> Optional[int]:
    for cell in cells:
        row_number = int(getattr(cell, "row", 0) or 0)
        if row_number:
            return row_number
    return None


def _row_has_value(cells, max_cell_chars: int) -> bool:
    for cell in cells:
        value, _truncated = _cell_value(cell, max_cell_chars)
        if value is not None:
            return True
    return False


def _row_values(
    cells,
    *,
    max_columns: int,
    cells_left: int,
    chars_left: int,
    max_cell_chars: int,
) -> tuple[list[Any], bool, int, int, bool]:
    values: list[Any] = []
    truncated = False
    extra_column = False
    for offset, cell in enumerate(cells, start=1):
        column = int(getattr(cell, "column", 0) or 0) or offset
        if column > max_columns:
            value, _char_truncated = _cell_value(cell, max_cell_chars)
            if value is not None:
                extra_column = True
            continue
        value, char_truncated = _cell_value(cell, max_cell_chars)
        truncated = truncated or char_truncated
        kept: Any = None
        output_clipped = False
        if value is not None:
            if cells_left <= 0 or chars_left <= 0:
                truncated = True
                break
            kept = value[0]
            if isinstance(kept, str) and len(kept) > chars_left:
                kept = kept[:chars_left]
                output_clipped = True
            elif _value_len(kept) > chars_left:
                truncated = True
                break
            chars_left -= _value_len(kept)
            cells_left -= 1
            truncated = truncated or output_clipped
        while len(values) < max(column - 1, 0):
            values.append(None)
        values.append(kept)
        if output_clipped:
            break
    return values, truncated, cells_left, chars_left, extra_column


def _value_len(value: Any) -> int:
    if isinstance(value, str):
        return len(value)
    if isinstance(value, bool):
        return 4 if value else 5
    return len(str(value))


def _cell_value(cell, max_cell_chars: int) -> tuple[Optional[tuple[Any, str]], bool]:
    value = getattr(cell, "value", None)
    if value is None:
        return None, False
    kind = getattr(cell, "data_type", None)
    if kind == "f":
        formula = getattr(value, "text", None)
        text_value = formula if isinstance(formula, str) and formula else str(value)
        text, truncated = _sanitize_text(text_value, max_cell_chars)
        if text and not text.startswith("="):
            text = "=" + text
        if not text:
            return None, truncated
        return (text, "formula"), truncated
    if isinstance(value, bool):
        return (value, "boolean"), False
    if isinstance(value, datetime):
        return (value.isoformat(), "date"), False
    if isinstance(value, date):
        return (value.isoformat(), "date"), False
    if isinstance(value, time):
        return (value.isoformat(), "time"), False
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return None, False
        return (value, "number"), False
    if isinstance(value, int):
        return (value, "number"), False
    text, truncated = _sanitize_text(str(value), max_cell_chars)
    if text == "":
        return None, truncated
    xml_type = "error" if kind == "e" else "string"
    return (text, xml_type), truncated


def _sanitize_text(text: str, max_cell_chars: int) -> tuple[str, bool]:
    cleaned = _ILLEGAL_XML.sub("", text)
    if len(cleaned) > max_cell_chars:
        return cleaned[:max_cell_chars], True
    return cleaned, False


def _trim_trailing(values: list[Any]) -> list[Any]:
    end = len(values)
    while end > 0 and values[end - 1] is None:
        end -= 1
    return values[:end]


def _sheet_document(
    name: str,
    parsed: list[tuple[int, list[Any]]],
    header_row: bool,
) -> dict[str, Any]:
    if not header_row:
        return {
            "name": name,
            "rows": [{"row": excel_row, "values": values} for excel_row, values in parsed],
        }
    if not parsed:
        return {"name": name, "header_row": None, "columns": [], "rows": []}
    header_excel_row, header_values = parsed[0]
    width = max(len(values) for _excel_row, values in parsed)
    if len(header_values) < width:
        header_values = header_values + [None] * (width - len(header_values))
    columns = _column_names(header_values)
    records = []
    for excel_row, values in parsed[1:]:
        record: dict[str, Any] = {_EXCEL_ROW_KEY: excel_row}
        for index, column in enumerate(columns):
            record[column] = values[index] if index < len(values) else None
        records.append(record)
    return {
        "name": name,
        "header_row": header_excel_row,
        "columns": columns,
        "rows": records,
    }


def _column_names(header_values: list[Any]) -> list[str]:
    used: set[str] = {_EXCEL_ROW_KEY}
    names: list[str] = []
    for index, value in enumerate(header_values, start=1):
        if isinstance(value, str) and value.strip():
            base = value.strip()
        elif value is None:
            base = get_column_letter(index)
        else:
            base = str(value).strip() or get_column_letter(index)
        candidate = base
        suffix = 2
        while candidate in used:
            candidate = f"{base}_{suffix}"
            suffix += 1
        used.add(candidate)
        names.append(candidate)
    return names


def _to_xml(document: dict[str, Any]) -> str:
    root = ET.Element(
        "workbook",
        {
            "file_name": str(document.get("file_name") or ""),
            "truncated": "true" if document.get("truncated") else "false",
        },
    )
    for sheet in document.get("sheets") or []:
        sheet_el = ET.SubElement(root, "sheet", {"name": str(sheet.get("name") or "")})
        columns = sheet.get("columns")
        if isinstance(columns, list):
            _append_record_rows(
                sheet_el,
                columns,
                sheet.get("rows") or [],
                sheet.get("header_row"),
            )
        else:
            _append_value_rows(sheet_el, sheet.get("rows") or [])
    body = ET.tostring(root, encoding="unicode")
    return '<?xml version="1.0" encoding="UTF-8"?>' + body


def _append_value_rows(sheet_el: ET.Element, rows: list[dict[str, Any]]) -> None:
    for record in rows:
        values = record.get("values") or []
        row_el = ET.SubElement(sheet_el, "row", {"index": str(record.get("row") or "")})
        _append_cells(row_el, values)


def _append_record_rows(
    sheet_el: ET.Element,
    columns: list[str],
    rows: list[dict[str, Any]],
    header_row: Any,
) -> None:
    header_attrs = {}
    if header_row is not None:
        header_attrs["index"] = str(header_row)
    header = ET.SubElement(sheet_el, "header", header_attrs)
    for index, column in enumerate(columns, start=1):
        cell = ET.SubElement(
            header,
            "cell",
            {"column": get_column_letter(index), "name": str(column)},
        )
        cell.text = str(column)
    for record in rows:
        row_el = ET.SubElement(
            sheet_el,
            "row",
            {"index": str(record.get(_EXCEL_ROW_KEY) or "")},
        )
        for index, column in enumerate(columns, start=1):
            _append_typed_cell(
                row_el,
                get_column_letter(index),
                record.get(column),
            )


def _append_cells(row_el: ET.Element, values: list[Any]) -> None:
    for index, value in enumerate(values, start=1):
        if value is None:
            continue
        _append_typed_cell(row_el, get_column_letter(index), value)


def _append_typed_cell(row_el: ET.Element, column: str, value: Any) -> None:
    if value is None:
        return
    cell = ET.SubElement(
        row_el,
        "cell",
        {"column": column, "type": _xml_type(value)},
    )
    cell.text = _xml_text(value)


def _xml_type(value: Any) -> str:
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return "number"
    if isinstance(value, str) and value.startswith("="):
        return "formula"
    return "string"


def _xml_text(value: Any) -> str:
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 10**15:
            return str(int(value))
        return format(value, ".16g")
    return str(value)
