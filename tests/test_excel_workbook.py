"""Excel transform tests for read_file. No network."""

import io
import unittest
import zipfile
from datetime import date

from openpyxl import Workbook

from app.core.exceptions import GraphValidationError
from app.services.excel_workbook import is_excel_file, transform_excel
from app.services.graph_service import GraphService, read_capped_chunks

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _orders_xlsx() -> bytes:
    workbook = Workbook()
    orders = workbook.active
    orders.title = "Orders"
    orders["A1"] = "Name"
    orders["B1"] = "Qty"
    orders["C1"] = "When"
    orders["A2"] = "Acme & Co <lab>"
    orders["B2"] = 3
    orders["C2"] = date(2026, 1, 15)
    orders["D2"] = "=B2*2"
    orders["B3"] = True
    orders["A5"] = "Name"
    notes = workbook.create_sheet("Notes")
    notes["A1"] = "only"
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


class ExcelTransformTest(unittest.TestCase):
    def test_json_preserves_values_dates_and_formulas(self):
        result = transform_excel(
            _orders_xlsx(),
            excel_format="JSON",
            file_name="orders.xlsx",
        )
        self.assertEqual(result["format"], "json")
        self.assertFalse(result["truncated"])
        self.assertIsNone(result["content_xml"])
        document = result["content_json"]
        self.assertEqual(document["file_name"], "orders.xlsx")
        self.assertEqual(document["truncated"], False)
        orders = document["sheets"][0]
        rows = {row["row"]: row["values"] for row in orders["rows"]}
        self.assertEqual(rows[1][:3], ["Name", "Qty", "When"])
        self.assertEqual(rows[2][0], "Acme & Co <lab>")
        self.assertEqual(rows[2][1], 3)
        self.assertTrue(str(rows[2][2]).startswith("2026-01-15"))
        self.assertEqual(rows[2][3], "=B2*2")
        self.assertNotEqual(rows[2][3], 6)
        self.assertEqual(rows[3], [None, True])
        self.assertEqual(rows[5], ["Name"])
        self.assertEqual(document["sheets"][1]["name"], "Notes")
        self.assertIn("Acme & Co <lab>", result["content_text"])

    def test_header_row_json_objects(self):
        document = transform_excel(
            _orders_xlsx(),
            excel_format="json",
            file_name="orders.xlsx",
            sheet_name="orders",
            header_row=True,
        )["content_json"]
        self.assertEqual(len(document["sheets"]), 1)
        sheet = document["sheets"][0]
        self.assertEqual(sheet["columns"][:4], ["Name", "Qty", "When", "D"])
        self.assertEqual(sheet["header_row"], 1)
        first = sheet["rows"][0]
        self.assertEqual(first["excel_row"], 2)
        self.assertEqual(first["Name"], "Acme & Co <lab>")
        self.assertEqual(first["Qty"], 3)
        self.assertEqual(first["D"], "=B2*2")
        self.assertEqual(sheet["rows"][1]["excel_row"], 3)
        self.assertTrue(sheet["rows"][1]["Qty"])
        self.assertEqual(sheet["rows"][2]["Name"], "Name")

    def test_duplicate_column_names_are_unique(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"] = "Name"
        sheet["B1"] = "Name"
        sheet["A2"] = "left"
        sheet["B2"] = "right"
        buffer = io.BytesIO()
        workbook.save(buffer)
        workbook.close()
        sheet_doc = transform_excel(
            buffer.getvalue(),
            excel_format="json",
            header_row=True,
        )["content_json"]["sheets"][0]
        self.assertEqual(sheet_doc["columns"], ["Name", "Name_2"])
        self.assertEqual(sheet_doc["rows"][0]["Name"], "left")
        self.assertEqual(sheet_doc["rows"][0]["Name_2"], "right")

    def test_xml_escapes_text(self):
        xml = transform_excel(
            _orders_xlsx(),
            excel_format="xml",
            file_name='orders"&.xlsx',
            sheet_name="Orders",
        )["content_xml"]
        self.assertTrue(xml.startswith("<?xml"))
        self.assertIn("Acme &amp; Co &lt;lab&gt;", xml)
        self.assertNotIn("<lab>", xml)
        self.assertIn("orders&quot;&amp;.xlsx", xml)
        self.assertIn('type="formula"', xml)
        self.assertIn("=B2*2", xml)
        self.assertIn('type="boolean"', xml)
        self.assertNotIn("<sheet name=\"Notes\"", xml)

    def test_truncates_extra_rows(self):
        result = transform_excel(
            _orders_xlsx(),
            excel_format="json",
            sheet_name="Orders",
            max_rows=1,
        )
        self.assertTrue(result["truncated"])
        rows = result["content_json"]["sheets"][0]["rows"]
        self.assertEqual([row["row"] for row in rows], [1])

    def test_truncates_long_cell_text(self):
        workbook = Workbook()
        workbook.active["A1"] = "abcdef"
        buffer = io.BytesIO()
        workbook.save(buffer)
        workbook.close()
        result = transform_excel(
            buffer.getvalue(),
            excel_format="json",
            max_cell_chars=3,
        )
        self.assertTrue(result["truncated"])
        self.assertEqual(result["content_json"]["sheets"][0]["rows"][0]["values"], ["abc"])

    def test_caps_total_output_characters(self):
        workbook = Workbook()
        workbook.active["A1"] = "abcdefghij"
        workbook.active["B1"] = "klmnopqrst"
        buffer = io.BytesIO()
        workbook.save(buffer)
        workbook.close()
        result = transform_excel(
            buffer.getvalue(),
            excel_format="json",
            max_output_chars=90,
        )
        self.assertTrue(result["truncated"])
        self.assertLessEqual(len(result["content_text"]), 90)
        self.assertNotIn("klmnopqrst", result["content_text"])

    def test_truncates_when_later_row_is_past_a_gap(self):
        result = transform_excel(
            _orders_xlsx(),
            excel_format="json",
            sheet_name="Orders",
            max_rows=3,
        )
        self.assertTrue(result["truncated"])
        rows = [row["row"] for row in result["content_json"]["sheets"][0]["rows"]]
        self.assertEqual(rows, [1, 2, 3])

    def test_truncates_when_later_column_is_past_a_gap(self):
        workbook = Workbook()
        workbook.active["A1"] = "keep"
        workbook.active["C1"] = "hidden"
        buffer = io.BytesIO()
        workbook.save(buffer)
        workbook.close()
        result = transform_excel(
            buffer.getvalue(),
            excel_format="json",
            max_columns=1,
        )
        self.assertTrue(result["truncated"])
        self.assertEqual(
            result["content_json"]["sheets"][0]["rows"][0]["values"],
            ["keep"],
        )

    def test_serialized_output_stays_within_cap(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"] = "H" * 40
        for index in range(2, 8):
            sheet.cell(index, 1, "x" * 40)
        buffer = io.BytesIO()
        workbook.save(buffer)
        workbook.close()
        result = transform_excel(
            buffer.getvalue(),
            excel_format="json",
            header_row=True,
            max_output_chars=120,
        )
        self.assertTrue(result["truncated"])
        self.assertLessEqual(len(result["content_text"]), 120)

    def test_defusedxml_parser_is_enabled(self):
        from openpyxl.xml import functions as xml_functions

        self.assertTrue(xml_functions.DEFUSEDXML)

    def test_rejects_non_deflate_compression(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as package:
            info = zipfile.ZipInfo("xl/workbook.xml")
            info.compress_type = zipfile.ZIP_BZIP2
            try:
                package.writestr(info, b"<workbook/>")
            except (NotImplementedError, RuntimeError):
                self.skipTest("bzip2 is not available")
        with self.assertRaises(GraphValidationError) as raised:
            transform_excel(buffer.getvalue(), excel_format="json")
        self.assertIn("Could not read", str(raised.exception))

    def test_strips_xml_illegal_characters(self):
        from app.services.excel_workbook import _sanitize_text

        cleaned, truncated = _sanitize_text("a\x00b\x0bc", 8000)
        self.assertEqual(cleaned, "abc")
        self.assertFalse(truncated)
        clipped, truncated = _sanitize_text("abcdef", 3)
        self.assertEqual(clipped, "abc")
        self.assertTrue(truncated)

        buffer = io.BytesIO()
        workbook = Workbook()
        workbook.active["A1"] = "ab"
        workbook.save(buffer)
        workbook.close()
        patched = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as source, zipfile.ZipFile(
            patched, "w"
        ) as target:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename.endswith(".xml"):
                    data = data.replace(b">ab<", b">a\x00b<")
                target.writestr(info, data)
        with self.assertRaises(GraphValidationError) as raised:
            transform_excel(patched.getvalue(), excel_format="json")
        self.assertNotIn("\x00", str(raised.exception))
        self.assertIn("Could not read", str(raised.exception))

    def test_missing_sheet(self):
        with self.assertRaises(GraphValidationError) as raised:
            transform_excel(_orders_xlsx(), excel_format="json", sheet_name="Missing")
        self.assertIn("Orders", str(raised.exception))
        self.assertIn("not found", str(raised.exception).lower())

    def test_legacy_xls_ole_is_rejected(self):
        with self.assertRaises(GraphValidationError) as raised:
            transform_excel(_OLE + b"legacy", excel_format="json", file_name="book.xls")
        self.assertIn(".xlsx", str(raised.exception))

    def test_encrypted_package_is_rejected(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as package:
            package.writestr("EncryptionInfo", "secret-marker")
            package.writestr("xl/workbook.xml", "<workbook/>")
        with self.assertRaises(GraphValidationError) as raised:
            transform_excel(buffer.getvalue(), excel_format="xml")
        self.assertIn("Encrypted", str(raised.exception))
        self.assertNotIn("secret-marker", str(raised.exception))

    def test_path_traversal_entry_is_rejected(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as package:
            package.writestr("../evil.xml", "<a/>")
            package.writestr("xl/workbook.xml", "<workbook/>")
        with self.assertRaises(GraphValidationError) as raised:
            transform_excel(buffer.getvalue(), excel_format="json")
        self.assertNotIn("evil", str(raised.exception))

    def test_size_caps(self):
        content = _orders_xlsx()
        with self.assertRaises(GraphValidationError):
            transform_excel(content, excel_format="json", max_bytes=32)
        with self.assertRaises(GraphValidationError):
            transform_excel(content, excel_format="json", max_uncompressed_bytes=32)
        with self.assertRaises(GraphValidationError):
            transform_excel(content, excel_format="json", max_zip_entries=1)

    def test_invalid_format(self):
        with self.assertRaises(GraphValidationError):
            transform_excel(_orders_xlsx(), excel_format="csv")

    def test_is_excel_file(self):
        self.assertTrue(is_excel_file("Book.XLSX", None))
        self.assertTrue(is_excel_file("book.xlsm", "application/octet-stream"))
        self.assertTrue(is_excel_file("download", _XLSX_MIME))
        self.assertFalse(is_excel_file("model.xlsb", None))
        self.assertFalse(is_excel_file("report.pdf", "application/pdf"))


class ReadFileExcelWiringTest(unittest.TestCase):
    def _service(self, content: bytes, item: dict) -> tuple[GraphService, list]:
        service = GraphService(credentials_json='{"access_token":"token-a"}')
        calls = []

        def fake_request(method, url, **kwargs):
            calls.append((method, url, kwargs.get("raw")))
            if kwargs.get("raw"):
                return content
            return item

        service._request = fake_request
        return service, calls

    def test_read_file_returns_json_without_base64(self):
        content = _orders_xlsx()
        item = {
            "id": "item-1",
            "name": "orders.xlsx",
            "size": len(content),
            "file": {"mimeType": _XLSX_MIME},
            "parentReference": {},
        }
        service, calls = self._service(content, item)
        result = service.read_file(item_id="item-1", excel_format="json", header_row=True)
        self.assertNotIn("content_base64", result)
        self.assertEqual(result["excel_format"], "json")
        self.assertEqual(result["content_json"]["sheets"][0]["columns"][0], "Name")
        self.assertEqual(result["metadata"]["name"], "orders.xlsx")
        self.assertEqual(len(calls), 2)

    def test_read_file_without_format_stays_base64(self):
        content = _orders_xlsx()
        item = {
            "id": "item-1",
            "name": "orders.xlsx",
            "size": len(content),
            "file": {"mimeType": _XLSX_MIME},
            "parentReference": {},
        }
        service, _calls = self._service(content, item)
        result = service.read_file(item_id="item-1")
        self.assertNotIn("content_json", result)
        self.assertTrue(result["content_base64"])
        raw, mime, name = service.download_file(item_id="item-1")
        self.assertEqual(raw, content)
        self.assertEqual(mime, _XLSX_MIME)
        self.assertEqual(name, "orders.xlsx")

    def test_rejects_non_excel_and_bad_format_without_download(self):
        content = b"%PDF-1.4"
        item = {
            "id": "item-2",
            "name": "report.pdf",
            "size": len(content),
            "file": {"mimeType": "application/pdf"},
            "parentReference": {},
        }
        service, calls = self._service(content, item)
        with self.assertRaises(GraphValidationError):
            service.read_file(item_id="item-2", excel_format="json")
        self.assertEqual(len(calls), 1)

        calls.clear()
        with self.assertRaises(GraphValidationError):
            service.read_file(item_id="item-2", excel_format="yaml")
        self.assertEqual(calls, [])

        with self.assertRaises(GraphValidationError):
            service.read_file(
                item_id="item-2",
                excel_format="json",
                include_content=False,
            )
        self.assertEqual(calls, [])


class DownloadCapTest(unittest.TestCase):
    def test_read_capped_chunks_stops_past_the_limit(self):
        self.assertEqual(read_capped_chunks([b"ab", b"cd"], 4), b"abcd")
        with self.assertRaises(GraphValidationError):
            read_capped_chunks([b"a" * 10, b"b" * 10], 12)

    def test_declared_size_rejects_before_content_download(self):
        item = {
            "id": "item-1",
            "name": "big.xlsx",
            "size": 8 * 1024 * 1024 + 1,
            "file": {"mimeType": _XLSX_MIME},
            "parentReference": {},
        }
        service = GraphService(credentials_json='{"access_token":"token-a"}')
        calls = []

        def fake_request(method, url, **kwargs):
            calls.append(kwargs.get("raw"))
            return item

        service._request = fake_request
        with self.assertRaises(GraphValidationError):
            service.read_file(item_id="item-1", excel_format="json")
        self.assertEqual(calls, [None])


if __name__ == "__main__":
    unittest.main()
